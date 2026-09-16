import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from report_worker.huaxing import HuaxingDownloader, write_roster
from report_worker.executor import RunExecutor, CancelledError
from converter import convert_file, filter_source_file
from target_login import TargetLogin


def config(tmp_path):
    return {"credentials": {"source_site": {"username": "fixture", "password": "fixture"}},
            "download": {"include_projects": ["测试工程"], "exclude_projects": []},
            "browser": {"download_dir": str(tmp_path)}}


def test_login_search_pagination_and_conversion(tmp_path, monkeypatch):
    source = HuaxingDownloader(config(tmp_path))
    requests = []
    def request(path, params=None, body=None):
        requests.append((path, params, body))
        if path == "/api/system/login":
            assert body == {"username": "fixture", "password": "fixture", "terminal": 1}
            return {"token": "temporary-test-token", "isResetPassword": True}
        if path == "/api/system/admin/self":
            assert source.token == "temporary-test-token"
            return {"user": {"id": 175}}
        if path == "/api/api/project/list":
            assert params["projectName"] == "测试工程"
            return {"lists": [{"id": 1474, "projectName": "测试工程"}, {"id": 1475, "projectName": "测试工程二期"}], "count": 2}
        assert params["projectId"] == 1474 and params["uid"] == 175 and params["workStatus"] == 0
        page = params["pageNo"]
        return {"count": 2, "lists": [{"id": page, "name": "测试工人", "idCard": "110101199001010011", "sex": 1,
                "phone": "13800000000", "address": "测试地址", "entryTime": str(date.today() - timedelta(days=0 if page == 1 else 10))}]}
    monkeypatch.setattr(source, "_request", request)
    files = source.run()
    assert len(files) == 1 and source.token is None
    assert len([r for r in requests if r[0] == "/api/api/worker/list"]) == 2
    assert filter_source_file(files[0], 1) == {"retained_count": 1, "filtered_count": 1}
    output = tmp_path / "converted"
    output.mkdir()
    result = convert_file(files[0], Path(__file__).resolve().parents[1] / "assets/建筑项目人员备案信息模板.xlsx", output, 1)
    wb = openpyxl.load_workbook(result)
    assert [wb['sheet1'].cell(3, c).value for c in (1, 2, 5, 6, 7)] == ["测试工人", "男", "110101199001010011", "13800000000", "测试地址"]
    wb.close()


def test_no_creation_timestamp_fallback(tmp_path):
    with pytest.raises(RuntimeError, match="进场时间"):
        write_roster(tmp_path / "source.xlsx", [{"id": 1, "createTime": str(date.today())}])


def test_repeated_page_rejected(tmp_path, monkeypatch):
    source = HuaxingDownloader(config(tmp_path))
    monkeypatch.setattr(source, "_request", lambda *args: {"count": 2, "lists": [{"id": 1}]})
    with pytest.raises(RuntimeError, match="重复"):
        list(source._pages("/api/api/worker/list", {}))


def test_formula_values_written_as_text(tmp_path):
    path = tmp_path / "source.xlsx"
    write_roster(path, [{"name": "=1+1", "entryTime": str(date.today())}])
    wb = openpyxl.load_workbook(path)
    assert wb['花名册']['B4'].data_type == 's'
    wb.close()


def test_manual_login_never_starts_feishu():
    target = TargetLogin({"credentials": {"target_site": {}}, "verification_type": "manual", "manual_code_provider": lambda **kw: "123456"})
    assert target._start_feishu_listener()
    assert target.listener_proc is None
    assert target._wait_for_sms_code() == "123456"


@pytest.mark.parametrize("cancelled", [False, True])
def test_manual_challenge_consumed_and_cleared(cancelled):
    class Repo:
        cleared = False
        def request_verification(self, run_id, timeout): return "request-1"
        def consume_verification(self, run_id, request_id): return "123456"
        def clear_verification(self, run_id): self.cleared = True
        def set_stage(self, *args): pass
        def event(self, *args): pass
        def cancelled(self, run_id): return cancelled
    executor = RunExecutor.__new__(RunExecutor)
    executor.repo = Repo()
    executor.run_id = "run-1"
    executor.context = {"stage": "target_login"}
    if cancelled:
        with pytest.raises(CancelledError): executor._manual_code()
    else:
        assert executor._manual_code() == "123456"
    assert executor.repo.cleared
    assert executor.context["stage"] == "target_login"


def test_manual_execution_failure_does_not_retry():
    class Repo:
        def event(self, *args): pass
        def complete(self, *args): self.completed = args
        def clear_verification(self, *args): pass
        def schedule_retry(self, *args): raise AssertionError("manual tasks cannot retry automatically")
    class Temp:
        def cleanup(self): pass
    executor = RunExecutor.__new__(RunExecutor)
    executor.repo = Repo()
    executor.run_id = "run"
    executor.config_row = {"verification_type": "manual"}
    executor.mode = "production"
    executor.context = {"stage": "download"}
    executor.temp = Temp()
    executor._download = lambda: (_ for _ in ()).throw(RuntimeError("network unavailable"))
    assert executor.execute() == "failed"
    assert executor.repo.completed[1] == "failed"


def test_local_worker_never_schedules_production():
    from report_worker.repository import Repository
    repo = Repository.__new__(Repository)
    repo.worker_target = "local"
    repo.connection = lambda: (_ for _ in ()).throw(AssertionError("local worker must not schedule"))
    repo.schedule_due()


def test_huaxing_file_names_survive_upload_splitting(tmp_path):
    from uploader import extract_project_name, split_upload_workbook
    path = tmp_path / "20260912_测试工程_华瑆导出.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "sheet1"
    for row in (3, 4):
        sheet.cell(row, 1, "测试工人")
        sheet.cell(row, 5, "110101199001010011")
    workbook.save(path)
    workbook.close()
    files = split_upload_workbook(path, 1)
    assert len(files) == 2
    assert all(extract_project_name(file) == "测试工程" for file in files)
    assert all("华瑆导出" in file for file in files)
